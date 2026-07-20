import zipfile
import xml.etree.ElementTree as ET
import os
import sys

def extract_docx_text(docx_path):
    if not os.path.exists(docx_path):
        print(f"File not found: {docx_path}")
        return
        
    try:
        import docx
        doc = docx.Document(docx_path)
        text = []
        for para in doc.paragraphs:
            text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text for cell in row.cells]
                text.append(" | ".join(row_text))
        return "\n".join(text)
    except ImportError:
        paragraphs = []
        namespaces = {
            'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        }
        with zipfile.ZipFile(docx_path) as z:
            xml_content = z.read('word/document.xml')
            root = ET.fromstring(xml_content)
            for p in root.findall('.//w:p', namespaces):
                p_text = []
                for t in p.findall('.//w:t', namespaces):
                    if t.text:
                        p_text.append(t.text)
                paragraphs.append("".join(p_text))
        return "\n".join(paragraphs)

if __name__ == "__main__":
    docx_path = r"d:\mockai\Temporal_Shift_Tabular_DL_Journal -sample for our reference.docx"
    text = extract_docx_text(docx_path)
    output_path = r"d:\mockai\scratch\sample_extracted.txt"
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(text)
    print(f"Extracted {len(text)} characters to {output_path}")
